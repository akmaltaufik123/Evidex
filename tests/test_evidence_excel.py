"""Regression tests for the Excel viewing copy."""
import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from evidence_excel import write_excel, COLUMNS


class ExcelTests(unittest.TestCase):
    def test_zeros_ids_formulas_unicode_and_missing_values(self):
        with tempfile.TemporaryDirectory() as directory:
            case = Path(directory)
            record = dict(message_id='00123', sender_id=1234567890123456789,
                          text='=HYPERLINK("https://example.invalid")', views=0,
                          forwards=0, replies=None, media={'size_bytes': 0})
            source = (json.dumps(record) + '\n').encode()
            (case / 'messages.jsonl').write_bytes(source)
            write_excel(case, 'Bukti saya 零 0')
            book = load_workbook(case / 'messages.xlsx')
            page = book['Messages']
            values = dict(zip(COLUMNS, [cell.value for cell in page[2]]))
            for key in ('views', 'forwards', 'media.size_bytes'):
                self.assertEqual(values[key], 0)
                cell = page.cell(2, COLUMNS.index(key) + 1)
                self.assertEqual(cell.data_type, 'n')
                self.assertEqual(cell.number_format, '0')
            self.assertTrue(page.sheet_view.showZeros)
            self.assertEqual(values['message_id'], '00123')
            self.assertEqual(values['sender_id'], '1234567890123456789')
            self.assertEqual(page.cell(2, COLUMNS.index('text') + 1).data_type, 's')
            self.assertEqual(values['text'], record['text'])
            self.assertIsNone(values['replies'])
            self.assertEqual(book['About']['B1'].value, 'Bukti saya 零 0')
            self.assertEqual((case / 'messages.jsonl').read_bytes(), source)
            book.close()

    def test_empty_export_has_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            case = Path(directory)
            (case / 'messages.jsonl').write_text('')
            write_excel(case, 'case')
            book = load_workbook(case / 'messages.xlsx')
            self.assertEqual(book['Messages'].max_row, 1)
            book.close()
