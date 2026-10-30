"""Excel viewing copy of captured JSONL; source evidence remains unchanged."""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font

COLUMNS = ('message_id', 'date_utc', 'edit_date_utc', 'sender_id', 'post_author',
           'text', 'reply_to_message_id', 'grouped_id', 'views', 'forwards', 'replies',
           'media.telegram_type', 'media.name', 'media.mime_type', 'media.size_bytes',
           'media.saved_path', 'media.sha256', 'media.download_error')
IDS = {'message_id', 'sender_id', 'reply_to_message_id', 'grouped_id'}


def write_excel(case: Path, name: str) -> None:
    """Stream a derived workbook, preserving zero values and literal strings."""
    book = Workbook(write_only=True)
    info = book.create_sheet('About')
    for row in [('Evidence name', name), ('Source', 'messages.jsonl'),
                ('Notes', 'Viewing copy. Missing values are blank; numeric zero is 0. IDs are text. '
                 'Unsupported Excel control characters are omitted only from this copy. '
                 'Keep JSONL and the original evidence files for verification.')]:
        cells = []
        for value in row:
            cell = WriteOnlyCell(info, value=ILLEGAL_CHARACTERS_RE.sub('', value))
            cell.data_type = 's'
            cells.append(cell)
        info.append(cells)

    def sheet(number):
        page = book.create_sheet('Messages' if number == 1 else f'Messages {number}')
        page.sheet_view.showZeros = True
        page.freeze_panes = 'A2'
        headers = []
        for key in COLUMNS:
            cell = WriteOnlyCell(page, value=key)
            cell.font = Font(bold=True)
            headers.append(cell)
        page.append(headers)
        return page

    number, rows = 1, 1
    page = sheet(number)
    try:
        with (case / 'messages.jsonl').open(encoding='utf-8') as source:
            for line in source:
                record = json.loads(line)
                if rows == 1048576:
                    page.auto_filter.ref = f'A1:R{rows}'
                    number += 1
                    page, rows = sheet(number), 1
                cells = []
                for key in COLUMNS:
                    value = record
                    for part in key.split('.'):
                        value = value.get(part) if isinstance(value, dict) else None
                    if value is not None and key in IDS:
                        value = str(value)
                    cell = WriteOnlyCell(page, value=value)
                    if isinstance(value, str):
                        cell.value = ILLEGAL_CHARACTERS_RE.sub('', value)
                        cell.data_type = 's'  # Telegram text must never become an Excel formula.
                        cell.number_format = '@'
                    elif isinstance(value, (int, float)):
                        cell.number_format = '0'
                    cells.append(cell)
                page.append(cells)
                rows += 1
        page.auto_filter.ref = f'A1:R{rows}'
        temporary = case / 'messages.xlsx.tmp'
        book.save(temporary)
        temporary.replace(case / 'messages.xlsx')
    finally:
        book.close()
